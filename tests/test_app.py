"""Testovi Pjesmarica konvertera.

Pokretanje:  pip install -r requirements.txt -r requirements-dev.txt && pytest -q
Testni PDF-ovi se generiraju reportlabom (nema vanjskih datoteka).
"""

import base64
import io
import json
import re
import zipfile

import pdfplumber
import pytest
from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app import (_skipped_msg, app, body_to_parts, editable_to_song,
                 parse_options, song_to_editable, strip_chord_lines)
from pdf_to_word import SECTION_LABELS, is_chord_line
from watermark import faded_logo_png, faded_png_from_bytes

W, H = A4
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# ── Pomoćne funkcije ──────────────────────────────────────────────────────────

def _song_page(c, title, key, capo, blocks):
    """Stranica nalik izvozu iz aplikacije s pjesmama (naslov, Key, sekcije)."""
    y = H - 60
    c.setFont("Helvetica-Bold", 24)
    c.drawString(50, y, title)
    c.setFont("Helvetica", 11)
    c.drawRightString(W - 50, y, key)
    y -= 16
    if capo:
        c.drawRightString(W - 50, y, capo)
    y -= 24
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, y, "Pjesmarica")
    y -= 14
    c.setFont("Helvetica", 11)
    c.drawString(50, y, "1 - 23")
    y -= 28
    for label, lines in blocks:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(50, y, label)
        y -= 14
        c.setFont("Helvetica", 11)
        for ln in lines:
            c.drawString(50, y, ln)
            y -= 14
        y -= 14
    c.showPage()


def _parse(client, pdf_bytes, name="Misa test.pdf"):
    return client.post("/parse", data={"pdf": (io.BytesIO(pdf_bytes), name)},
                       content_type="multipart/form-data")


def _export(client, kind, songs, opts=None, **extra):
    body = {"filename": "Misa", "options": opts or {}, "songs": songs}
    body.update(extra)
    return client.post(f"/export/{kind}", data=json.dumps(body),
                       content_type="application/json")


def _sect_pr(docx_bytes):
    doc_xml = zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml").decode()
    return re.search(r"<w:sectPr.*?</w:sectPr>", doc_xml, re.S).group(0)


# ── Fixturi ───────────────────────────────────────────────────────────────────

@pytest.fixture(scope="module")
def songbook_pdf():
    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    _song_page(c, "Radujte se narodi", "Key: A", "Capo 1", [
        ("Verse 1", ["D A Fis H",
                     "Radujte se narodi kad cujete glas,",
                     "da se Isus rodio za sve nas"]),
        ("Chorus", ["Aleluja, aleluja,", "rodio se Spasitelj svijeta"]),
    ])
    _song_page(c, "Druga pjesma", "Key: G", "", [
        ("Verse 1", ["Prvi redak druge pjesme,", "drugi redak druge pjesme"]),
        ("Bridge", ["Most pjesme ide ovdje"]),
    ])
    c.save()
    return b.getvalue()


@pytest.fixture
def client():
    return app.test_client()


EDITED = [
    {"title": "Radujte se narodi", "key": "Key: A\nCapo 1", "subtitle": "",
     "pjesmarica": "1 - 23",
     "body": "Verse 1\nD A Fis H\nRadujte se narodi kad čujete glas,\n"
             "da se Isus rodio za sve nas\nChorus\nAleluja, aleluja (ćčžšđ)"},
    {"title": "Druga pjesma", "key": "Key: G", "subtitle": "Papa Band",
     "pjesmarica": "",
     "body": "Verse 1\nPrvi redak s čćžšđ\nBridge\nMost pjesme"},
]


# ── Stranica, verzija, PWA ────────────────────────────────────────────────────

def test_index_no_cache_and_version(client):
    r = client.get("/")
    html = r.get_data(as_text=True)
    assert r.status_code == 200
    assert "no-cache" in r.headers["Cache-Control"]
    assert "Pjesmarica konverter" in html
    assert "__VERSION__" not in html and "__MAX_MB__" not in html
    assert 'fetch("parse"' in html and 'fetch("export/"' in html  # relativne rute


def test_health_has_version(client):
    h = client.get("/health").get_json()
    assert h["status"] == "ok" and h["version"]


def test_pwa_manifest_sw_icons(client):
    r = client.get("/manifest.webmanifest")
    m = r.get_json()
    assert r.status_code == 200 and r.mimetype == "application/manifest+json"
    assert m["start_url"] == "./" and m["display"] == "standalone"
    assert all(not i["src"].startswith("/") for i in m["icons"])
    for icon in m["icons"]:
        assert client.get("/" + icon["src"]).mimetype == "image/png"
    r = client.get("/sw.js")
    assert "javascript" in r.mimetype and "no-cache" in r.headers["Cache-Control"]
    assert "caches" not in r.get_data(as_text=True)  # SW ne kešira
    html = client.get("/").get_data(as_text=True)
    assert 'rel="manifest"' in html and "apple-touch-icon" in html


# ── Parser i pomoćne funkcije ─────────────────────────────────────────────────

def test_chord_detection():
    for ln in ("D A Fis H", "Am7 G/B Dsus4", "C", "Es Bb", "D A Hm G A D"):
        assert is_chord_line(ln), ln
    for ln in ("Radujte se narodi", "A on je nas", "Da te samo dotaknem", "Amen", ""):
        assert not is_chord_line(ln), ln


def test_section_labels():
    for ln in ("Verse 1", "Chorus", "Chorus 2", "Bridge", "V2 (tiho)", "C1", "Pre-Chorus"):
        assert SECTION_LABELS.match(ln), ln
    assert not SECTION_LABELS.match("Radujte se narodi")


def test_parse_songbook(client, songbook_pdf):
    d = _parse(client, songbook_pdf, "Misa_7192026_print.pdf").get_json()
    assert d["filename"] == "Misa_7192026"          # sufiks _print skinut
    assert len(d["songs"]) == 2 and d["pages"] == 2 and d["skipped"] == 0
    assert d["warnings"] == []
    s0 = d["songs"][0]
    assert s0["title"] == "Radujte se narodi"
    assert "Key: A" in s0["key"] and "Capo 1" in s0["key"]
    assert s0["pjesmarica"] == "1 - 23"
    assert "D A Fis H" in s0["body"]                 # akordi vidljivi u editoru
    assert "Verse 1" in s0["body"] and "Chorus" in s0["body"]


def test_roundtrip_editable(client, songbook_pdf):
    s0 = _parse(client, songbook_pdf).get_json()["songs"][0]
    song = editable_to_song(s0)
    assert song["sections"][0][0] == "Verse 1" and song["sections"][1][0] == "Chorus"
    assert song_to_editable(song)["body"] == s0["body"]


def test_body_to_parts_and_strip_chords():
    notes, sections = body_to_parts("Napomena\n\nVerse 1\nD A\nredak 1\n\nredak 2\nChorus\nref")
    assert notes == ["Napomena"]
    assert sections[0] == ["Verse 1", ["D A", "redak 1", "", "redak 2"]]
    song = strip_chord_lines({"notes": notes, "sections": sections})
    assert song["sections"][0][1] == ["redak 1", "", "redak 2"]


def test_parse_options_defaults_and_clamp():
    o = parse_options({})
    assert o["layout"] == "kompaktno" and o["font"] == "Liberation Serif"
    assert o["body_size"] == 9 and o["n_cols"] == 3 and o["strip_chords"] is True
    assert o["watermark"] is False and o["watermark_level"] == "vrlo_blijedo"
    o = parse_options({"layout": "klasicno", "font": "Comic Sans", "n_cols": 7,
                       "body_size": 1, "strip_chords": False,
                       "watermark": True, "watermark_level": "x"})
    assert o["layout"] == "klasicno" and o["font"] == "Liberation Serif"
    assert o["n_cols"] == 3 and o["body_size"] == 7 and o["strip_chords"] is False
    assert o["watermark"] is True and o["watermark_level"] == "vrlo_blijedo"


def test_parse_scan_without_text(client):
    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    c.showPage()
    c.showPage()
    c.save()
    r = _parse(client, b.getvalue(), "sken.pdf")
    assert r.status_code == 422 and "sken" in r.get_json()["error"]


def test_parse_skipped_page_warning(client, songbook_pdf):
    b = io.BytesIO()
    c = canvas.Canvas(b, pagesize=A4)
    _song_page(c, "Prva", "Key: A", "", [("Verse 1", ["redak"])])
    c.setFont("Helvetica", 11)
    c.drawString(420, H - 100, "samo desno")  # tekst bez naslova lijevo
    c.showPage()
    c.save()
    d = _parse(client, b.getvalue()).get_json()
    assert len(d["songs"]) == 1 and d["pages"] == 2 and d["skipped"] == 1
    assert any("Preskočena 1 stranica od 2" in w for w in d["warnings"])


def test_skipped_msg_plurals():
    assert "Preskočena 1 stranica" in _skipped_msg(1, 12)
    assert "Preskočene 2 stranice" in _skipped_msg(2, 12)
    assert "Preskočeno 5 stranica" in _skipped_msg(5, 12)
    assert "Preskočeno 11 stranica" in _skipped_msg(11, 20)
    assert "Preskočena 21 stranica" in _skipped_msg(21, 30)
    assert "Preskočene 22 stranice" in _skipped_msg(22, 30)


def test_parse_already_converted_input_warns(client):
    converted = _export(client, "pdf", EDITED).data          # ležeći ispis s crtama
    d = _parse(client, converted, "Misa_print.pdf").get_json()
    assert any("KONVERTIRANI" in w for w in d["warnings"])


def test_parse_errors(client):
    assert client.post("/parse", data={}, content_type="multipart/form-data").status_code == 400
    r = client.post("/parse", data={"pdf": (io.BytesIO(b"x"), "a.txt")},
                    content_type="multipart/form-data")
    assert r.status_code == 400
    r = client.post("/parse", data={"pdf": (io.BytesIO(b"nije pdf"), "los.pdf")},
                    content_type="multipart/form-data")
    assert r.status_code == 500 and "Ne mogu" in r.get_json()["error"]


# ── Izvoz: kompaktni (original) ───────────────────────────────────────────────

def test_export_compact_pdf(client):
    r = _export(client, "pdf", EDITED, {"n_cols": 3})
    assert r.status_code == 200 and r.data[:5] == b"%PDF-"
    assert "Misa_print.pdf" in r.headers["Content-Disposition"]
    with pdfplumber.open(io.BytesIO(r.data)) as pdf:
        page = pdf.pages[0]
        assert page.width > page.height                    # ležeće
        text = page.extract_text() or ""
    assert "Radujte se narodi kad" in text and "čujete" in text and "ćčžšđ" in text
    assert "D A Fis H" not in text                         # akordi uklonjeni
    assert "Key: A" not in text                            # bez tonaliteta
    assert "_____" in text                                 # crta između pjesama


def test_export_compact_pdf_keeps_chords_when_disabled(client):
    r = _export(client, "pdf", EDITED, {"strip_chords": False})
    with pdfplumber.open(io.BytesIO(r.data)) as pdf:
        assert "D A Fis H" in (pdf.pages[0].extract_text() or "")


def test_export_compact_docx(client):
    r = _export(client, "docx", EDITED, {"n_cols": 3})
    assert r.status_code == 200
    assert "Misa_print.docx" in r.headers["Content-Disposition"]
    sect = _sect_pr(r.data)
    order = re.findall(r"<w:(\w+)", sect)[1:]
    assert sect.count("<w:cols") == 1                      # nema duplikata
    assert order.index("cols") < order.index("docGrid")   # ispravan redoslijed
    assert re.search(r'<w:cols [^>]*w:num="3"', sect)
    assert 'w:orient="landscape"' in sect
    doc = Document(io.BytesIO(r.data))
    sec = doc.sections[0]
    assert abs(sec.page_width.cm - 29.7) < 0.01 and abs(sec.page_height.cm - 21.0) < 0.01
    paras = [p for p in doc.paragraphs if p.text.strip()]
    texts = [p.text for p in paras]
    assert "Radujte se narodi" not in texts                # bez naslova
    assert not any("D A Fis H" in t for t in texts)        # bez akorda
    chorus = next(p for p in paras if "Aleluja" in p.text)
    assert chorus.runs[0].font.bold and chorus.runs[0].font.italic
    assert chorus.paragraph_format.left_indent is not None
    assert any("____" in t for t in texts)
    assert paras[0].runs[0].font.name == "Liberation Serif"


# ── Izvoz: klasični ───────────────────────────────────────────────────────────

def test_export_classic_pdf_and_docx(client):
    opts = {"layout": "klasicno", "font": "Liberation Sans", "title_size": 30,
            "body_size": 12, "margin_cm": 2.0, "page_break": True}
    r = _export(client, "pdf", EDITED, opts)
    assert r.status_code == 200
    with pdfplumber.open(io.BytesIO(r.data)) as pdf:
        assert len(pdf.pages) == 2
        assert pdf.pages[0].width < pdf.pages[0].height    # uspravno
        p1 = pdf.pages[0].extract_text() or ""
    assert "Radujte se narodi" in p1 and "Key: A" in p1 and "Capo 1" in p1
    r = _export(client, "docx", EDITED, opts)
    doc = Document(io.BytesIO(r.data))
    run = doc.paragraphs[0].runs[0]
    assert run.font.name == "Liberation Sans" and run.font.size.pt == 30
    sect = _sect_pr(r.data)
    assert sect.count("<w:cols") == 1 and 'w:orient="landscape"' not in sect


# ── Validacija ulaza ──────────────────────────────────────────────────────────

def test_export_validation(client):
    for bad in (["nije dict"], "x", [1, 2]):
        r = _export(client, "pdf", bad)
        assert r.status_code == 400 and "popis pjesama" in r.get_json()["error"]
    assert _export(client, "pdf", []).status_code == 400
    r = client.post("/export/pdf", data="nije json", content_type="text/plain")
    assert r.status_code == 400
    r = client.post("/export/pdf", data=json.dumps([1, 2]), content_type="application/json")
    assert r.status_code == 400


# ── Vodeni žig ────────────────────────────────────────────────────────────────

def test_watermark_pdf_and_docx(client):
    for layout in ("kompaktno", "klasicno"):
        off = _export(client, "pdf", EDITED, {"layout": layout, "watermark": False})
        on = _export(client, "pdf", EDITED, {"layout": layout, "watermark": True})
        assert on.status_code == 200 and len(on.data) > len(off.data) + 5000
        r = _export(client, "docx", EDITED, {"layout": layout, "watermark": True})
        z = zipfile.ZipFile(io.BytesIO(r.data))
        names = z.namelist()
        hdr = [n for n in names if n.startswith("word/header") and n.endswith(".xml")]
        assert hdr and "v:shape" in z.read(hdr[0]).decode()
        assert [n for n in names if n.startswith("word/media/")]
        assert "headerReference" in z.read("word/document.xml").decode()
    r = _export(client, "docx", EDITED, {"watermark": False})
    assert not [n for n in zipfile.ZipFile(io.BytesIO(r.data)).namelist()
                if n.startswith("word/media/")]


def test_faded_logo_levels():
    a = Image.open(io.BytesIO(faded_logo_png("vrlo_blijedo")))
    b = Image.open(io.BytesIO(faded_logo_png("srednje")))
    assert a.size == b.size and a.mode == "RGBA"
    assert a.getchannel("A").getextrema()[1] < b.getchannel("A").getextrema()[1]


def test_custom_logo_white_removed_and_cropped(client):
    im = Image.new("RGB", (1000, 600), "white")
    ImageDraw.Draw(im).rectangle([600, 100, 900, 300], fill="black")
    buf = io.BytesIO()
    im.save(buf, "PNG")
    out = Image.open(io.BytesIO(faded_png_from_bytes(buf.getvalue(), "srednje")))
    assert out.size == (301, 201)                            # obrezano na sadržaj
    assert out.getpixel((150, 100))[3] > 0                   # linija ostaje
    logo_url = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    custom = _export(client, "pdf", EDITED, {"watermark": True}, logo=logo_url)
    default = _export(client, "pdf", EDITED, {"watermark": True})
    assert custom.status_code == 200 and custom.data != default.data
    bad = _export(client, "pdf", EDITED, {"watermark": True},
                  logo="data:image/png;base64,bm90YW5pbWFnZQ==")
    assert bad.status_code == 200                            # fallback na ugrađeni
